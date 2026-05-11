"""从 docx 中抽取所有 "" 引文（不再要求段落带高亮）。

实际写作中研究者一般不会把引用段落标黄，所以扫描全部段落。
噪声（对话、强调引号等）由匹配阶段自然过滤掉——找不到对应书页的就标"未命中"。
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Union

from docx import Document


# 兼容中文弯引号 / 直角引号 / 英文直引号
QUOTE_PATTERN = re.compile(r"[“「\"]([^“”「」\"]+?)[”」\"]")

# 太短的引文价值低又容易误命中（如对话里的"哦""嗯"），过滤掉
_MIN_QUOTE_LEN = 3


@dataclass
class Quote:
    quote_id: int
    text: str
    paragraph_index: int
    paragraph_text: str
    char_start: int
    char_end: int

    @property
    def context_before(self) -> str:
        return self.paragraph_text[max(0, self.char_start - 30): self.char_start]

    @property
    def context_after(self) -> str:
        return self.paragraph_text[self.char_end: self.char_end + 30]


def extract_quotes(source: Union[Path, str, bytes, bytearray]) -> List[Quote]:
    """接受 Path（桌面版兼容）或 bytes（网页版）。"""
    if isinstance(source, (bytes, bytearray)):
        doc = Document(io.BytesIO(bytes(source)))
    else:
        doc = Document(str(source))
    quotes: List[Quote] = []
    qid = 0
    for p_idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text
        for m in QUOTE_PATTERN.finditer(text):
            inner = m.group(1).strip()
            if len(inner) < _MIN_QUOTE_LEN:
                continue
            qid += 1
            quotes.append(
                Quote(
                    quote_id=qid,
                    text=inner,
                    paragraph_index=p_idx,
                    paragraph_text=text,
                    char_start=m.start(1),
                    char_end=m.end(1),
                )
            )
    return quotes


if __name__ == "__main__":
    import sys

    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("ai查询并标注引文.docx")
    sys.stdout.reconfigure(encoding="utf-8")
    qs = extract_quotes(path)
    print(f"共抽到 {len(qs)} 条引文：\n")
    for q in qs:
        print(f"[{q.quote_id}] {q.text}")
        print(f"    ……{q.context_before}「{q.text}」{q.context_after}……")
